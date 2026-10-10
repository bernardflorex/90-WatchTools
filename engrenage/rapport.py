"""Rapport de simulation (.docx) : conditions d'optimisation, résultats sur la plage d'entraxe
et graphiques d'engrènement, pour un profil en arcs de cercle (menant/mené).

Les résultats chiffrés viennent tous de `contact.Engrenement` (le moteur analytique arc-à-arc,
validé par l'utilisateur contre le programme Fortran), jamais de `optim.py` (qui sert seulement
à retrouver les conditions d'optimisation déjà enregistrées dans `evolvente.Projet.optimisation`).
Mise en page avec `python-docx` (bibliothèque Python pure, pas de dépendance Node.js, pour rester
utilisable dans l'application une fois celle-ci livrée).
"""

import math
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import contact
import evolvente

N_POS = 150  # positions simulées par entraxe, comme NJ=0 du Fortran (IK=101, mais 150 pour l'affichage)
N_JEU = 101  # positions pour le jeu (assez pour un résultat stable, moins coûteux que 150)
COULEUR_TITRE = RGBColor(0x1F, 0x3A, 0x5F)
CRITERES_LABELS = {"couple": "variation de couple", "jeu": "jeu maximal", "rendement": "rendement",
                   "hertz": "pression de Hertz"}


def plage_da(da_neg, da_pos, pas_um=10.0):
    """ΔA (µm) de `da_neg` à `da_pos`, bornes incluses, par pas d'environ `pas_um`."""
    n = max(1, round((da_pos - da_neg) / pas_um))
    return [da_neg + i * (da_pos - da_neg) / n for i in range(n + 1)]


def simuler_plage(roues, profils, mobiles, mu, da_neg, da_pos, pas_um=10.0, n_pos=N_POS, n_jeu=N_JEU):
    """Simulation complète (comme l'onglet Simulation) à chaque entraxe de la plage.

    Retourne une liste de dicts, un par ΔA : `da_um`, `da_mm`, et soit `sim`
    (contact.Simulation) + `stats` (statistiques()) + `jeu_min`/`jeu_max`/`hertz_max`, soit
    `erreur` (str) si cet entraxe n'admet pas d'engrènement (`contact.ErreurContact`).
    """
    lignes = []
    for da_um in plage_da(da_neg, da_pos, pas_um):
        da_mm = da_um / 1000.0
        ligne = {"da_um": da_um, "da_mm": da_mm}
        try:
            moteur = contact.Engrenement(roues, profils, da_mm, mobiles=mobiles)
            sim = moteur.simuler(mu, n_pos, n_jeu=n_jeu)
            hz, _ = sim.hertz_max()
            ligne.update(sim=sim, stats=sim.statistiques(), jeu_min=sim.jeu_min, jeu_max=sim.jeu_max,
                        hertz_max=hz)
        except (contact.ErreurContact, ValueError, ArithmeticError) as e:
            ligne["erreur"] = str(e)
        lignes.append(ligne)
    return lignes


def resume(lignes):
    """Extrêmes des critères sur toute la plage (pour le résumé), None si aucune ligne valide."""
    ok = [l for l in lignes if "stats" in l]
    if not ok:
        return None
    return {
        "varC_max": max(l["stats"]["couple"][3] for l in ok),
        "rendement_min": min(l["stats"]["rendement"][0] for l in ok),
        "hertz_max": max(l["hertz_max"] for l in ok if l["hertz_max"] == l["hertz_max"]),
        "jeu_min": min(l["jeu_min"] for l in ok),
        "jeu_max": max(l["jeu_max"] for l in ok),
        "n_ok": len(ok), "n_total": len(lignes),
    }


# --------------------------------------------------------------------------------------------
# graphiques
# --------------------------------------------------------------------------------------------

def figure_engrenement(chemin_png, roues, profils, mobiles, sim, da_mm, titre):
    """Une image « engrènement d'une paire de dents » (comme l'onglet Simulation), position
    médiane de la période, avec les courbes W/C/N superposées."""
    i = len(sim) // 2
    positions = [sim.position(j) for j in (0, i, len(sim) - 1)]
    superposition = (sim.contact[:, 0].tolist(), {k: v.tolist() for k, v in sim.pourcent().items()}, i)
    limites = evolvente.limites_engrenage(roues, profils, 3, positions,
                                          evolvente.points_courbes(superposition, roues[0].mo), da_mm)
    fig = plt.figure(figsize=(7.2, 6.4), dpi=150)
    ax = fig.add_subplot(111)
    evolvente.dessiner_engrenage(ax, roues, profils, 3, positions[1], sim.contact.tolist(), limites,
                                 superposition, da_mm, None, [(m.points, m.normales()) for m in mobiles],
                                 sim.paire.tolist(), sim.hertz_local)
    ax.set_title(titre, fontsize=11)
    fig.tight_layout()
    fig.savefig(chemin_png)
    plt.close(fig)


def figure_tendance(chemin_png, lignes, jeu_min_admis=None, jeu_max_admis=None):
    """Variation de couple, pression de Hertz et jeu de la menée en fonction de l'entraxe."""
    ok = [l for l in lignes if "stats" in l]
    das = [l["da_um"] for l in ok]
    varc = [l["stats"]["couple"][3] for l in ok]
    hertz = [l["hertz_max"] for l in ok]
    jmin, jmax = [l["jeu_min"] for l in ok], [l["jeu_max"] for l in ok]
    fig, axes = plt.subplots(3, 1, figsize=(7.2, 6.8), dpi=150, sharex=True)
    axes[0].plot(das, varc, "o-", color="tab:brown", ms=4)
    axes[0].set_ylabel("Variation de\ncouple (%)")
    axes[1].plot(das, hertz, "o-", color="tab:red", ms=4)
    axes[1].set_ylabel("Hertz max\n(N/mm², 1 mN·m)")
    axes[2].fill_between(das, jmin, jmax, alpha=0.2, color="tab:blue", label="plage sur la période")
    axes[2].plot(das, jmin, "o-", color="tab:blue", ms=4, label="jeu minimal")
    axes[2].plot(das, jmax, "s-", color="tab:blue", ms=4, mfc="white", label="jeu maximal")
    if jeu_min_admis is not None:
        axes[2].axhline(jeu_min_admis, color="tab:red", lw=1, ls="--", label="jeu min. admissible")
    if jeu_max_admis:
        axes[2].axhline(jeu_max_admis, color="tab:red", lw=1, ls=":", label="jeu max. admissible")
    axes[2].set_ylabel("Jeu de la\nmenée (°)")
    axes[2].set_xlabel("Variation d'entraxe ΔA (µm)")
    axes[2].legend(fontsize="x-small", loc="best")
    for ax in axes:
        ax.grid(alpha=0.3)
        ax.axvline(0, color="k", lw=0.6)
    fig.suptitle("Critères en fonction de l'entraxe", fontsize=12)
    fig.tight_layout()
    fig.savefig(chemin_png)
    plt.close(fig)


# --------------------------------------------------------------------------------------------
# mise en page .docx
# --------------------------------------------------------------------------------------------

def _titre(doc, texte, niveau=1):
    h = doc.add_heading(texte, level=niveau)
    for run in h.runs:
        run.font.color.rgb = COULEUR_TITRE
    return h


def _tableau(doc, entetes, lignes, largeurs_cm=None):
    """Tableau avec en-tête grisé ; `lignes` : liste de listes de textes (déjà formatés)."""
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = largeurs_cm is None
    for j, texte in enumerate(entetes):
        cell = t.rows[0].cells[j]
        cell.text = texte
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, ligne in enumerate(lignes, start=1):
        for j, texte in enumerate(ligne):
            cell = t.rows[i].cells[j]
            cell.text = str(texte)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if largeurs_cm:
        for row in t.rows:
            for j, larg in enumerate(largeurs_cm):
                row.cells[j].width = Cm(larg)
    return t


def _libelles_libres(libres, n_arcs):
    if libres is None:
        return "tous"
    return ", ".join(str(i + 1) for i in sorted(libres)) or "aucun"


def ecrire_rapport(chemin, roues, mobiles, mu, mt, da_neg, da_pos, optimisation=None, origine="",
                   pas_um=10.0, dossier_images=None):
    """Écrit le rapport de simulation à `chemin` (.docx).

    `roues` : les deux `evolvente.Roue` (module, nombre de dents — les autres champs ne sont
    pas utilisés, comme dans l'onglet Simulation). `mobiles` : les deux `arcs.Mobile` du profil
    à documenter. `optimisation` : `evolvente.Projet.optimisation` si disponible (sinon la
    section "Conditions d'optimisation" est omise). `origine` : d'où vient ce profil (fichier,
    onglet…), pour la page de titre.
    """
    dossier_images = Path(dossier_images) if dossier_images else Path(chemin).with_suffix("")
    dossier_images.mkdir(parents=True, exist_ok=True)
    profils = [m.echantillonner().tolist() for m in mobiles]
    lignes = simuler_plage(roues, profils, mobiles, mu, da_neg, da_pos, pas_um)
    res = resume(lignes)
    m1, m2 = mobiles

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.0)
        section.left_margin = section.right_margin = Cm(2.2)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)

    # -- page de titre -------------------------------------------------------------------
    titre = doc.add_heading("Rapport de simulation d'engrènement", level=0)
    for run in titre.runs:
        run.font.color.rgb = COULEUR_TITRE
    p = doc.add_paragraph()
    p.add_run(f"{m1.nom} (menante, Z={m1.zz}) / {m2.nom} (menée, Z={m2.zz})  —  module {m1.mo:g} mm\n"
              f"{len(m1.points) - 1} / {len(m2.points) - 1} arcs de cercle"
              + ("  —  dents dissymétriques" if not (m1.symetrique and m2.symetrique) else "")).italic = True
    doc.add_paragraph(f"Origine du profil : {origine or 'non précisée'}").runs[0].italic = True
    doc.add_paragraph(f"Établi le {date.today().strftime('%d.%m.%Y')} — moteur de simulation "
                      "contact.py (analytique, arcs de cercle), validé contre le programme "
                      "Fortran ENGCRE.").runs[0].italic = True

    # -- 1. profil et conditions ----------------------------------------------------------
    _titre(doc, "1. Profil et conditions de simulation")
    _tableau(doc, ["", "Menante", "Menée"], [
        ["Désignation", m1.nom, m2.nom], ["Nombre de dents Z", m1.zz, m2.zz],
        ["Module (mm)", f"{m1.mo:g}", f"{m2.mo:g}"], ["Arcs de cercle", len(m1.points) - 1, len(m2.points) - 1],
        ["Module d'Young (N/mm²)", f"{m1.e_mod:g}" if m1.e_mod else "-", f"{m2.e_mod:g}" if m2.e_mod else "-"],
        ["Coeff. de Poisson", f"{m1.poisson:g}" if m1.poisson else "-", f"{m2.poisson:g}" if m2.poisson else "-"],
        ["Épaisseur (mm)", f"{m1.epaisseur:g}" if m1.epaisseur else "-", f"{m2.epaisseur:g}" if m2.epaisseur else "-"],
    ], largeurs_cm=[5.5, 5.0, 5.0])
    doc.add_paragraph()
    doc.add_paragraph(
        f"Coefficient de frottement μ = {mu:g}   —   couple sur la menante Cm = {mt:g} mN·m   —   "
        f"plage d'entraxe étudiée : {da_neg:+g} à {da_pos:+g} µm, par pas d'environ {pas_um:g} µm "
        f"({len(lignes)} positions).")
    doc.add_paragraph(
        "Les pressions de Hertz de ce rapport sont toujours données pour un couple moteur de "
        "référence de 1 mN·m, comme dans l'onglet Simulation ; pour un autre couple Cm, la "
        "pression réelle vaut la valeur du tableau multipliée par √Cm (Cm en mN·m).").runs[0].italic = True

    # -- 2. conditions d'optimisation -------------------------------------------------------
    if optimisation:
        _titre(doc, "2. Conditions d'optimisation")
        libelles = {cle: lib for cle, lib, _ in (("couple", "Variation de couple mini", None),
                                                  ("jeu", "Jeu mini", None), ("rendement", "Rendement maxi", None),
                                                  ("hertz", "Pression de Hertz mini", None))}
        princ = libelles.get(optimisation.get("principal"), optimisation.get("principal", "-"))
        sec = libelles.get(optimisation.get("secondaire"), optimisation.get("secondaire") or "aucun")
        doc.add_paragraph(f"Critère principal : {princ}" + (f"   —   secondaire : {sec}" if sec != "aucun" else "")
                          + (f"   (poids {optimisation.get('ponderation', 0):g})" if sec != "aucun" else "") + ".")
        _tableau(doc, ["Contrainte", "Valeur"], [
            ["Jeu minimal admissible (entraxe mini)", f"{optimisation.get('jeu_min_admis', 0):g} °"],
            ["Jeu maximal admissible (entraxe maxi)",
             f"{optimisation.get('jeu_max_admis', 0):g} °" if optimisation.get("jeu_max_admis") else "profil de départ + 0,02°"],
            ["Rayon minimal de l'arc de tête", f"{optimisation.get('r_tete_min', 0):g} mm" if optimisation.get("r_tete_min") else "aucun"],
            ["Interférence entre les dents", "toujours interdite (contours complets, tous entraxes)"],
            ["Réparation par demi-profil arrière", "oui" if optimisation.get("dissymetrie", 1) else "non"],
        ], largeurs_cm=[9.5, 6.0])
        doc.add_paragraph()
        libres = optimisation.get("libres", {})
        _tableau(doc, ["", "Arcs libres"], [
            ["Menante", _libelles_libres(libres.get("m"), len(m1.points) - 2)],
            ["Menée", _libelles_libres(libres.get("e"), len(m2.points) - 2)],
        ], largeurs_cm=[5.5, 10.0])
        doc.add_paragraph()
        doc.add_paragraph(f"Recherche : population {optimisation.get('population', '-')}, "
                          f"{optimisation.get('evaluations', '-')} évaluations par départ, "
                          f"{optimisation.get('departs', '-')} départs, graine {optimisation.get('graine', '-')}.")

    # -- 3. résumé ---------------------------------------------------------------------
    _titre(doc, "3. Résumé sur la plage d'entraxe")
    if res is None:
        doc.add_paragraph("Aucun entraxe de la plage n'admet d'engrènement : voir le tableau détaillé.")
    else:
        if res["n_ok"] < res["n_total"]:
            doc.add_paragraph(f"⚠ {res['n_total'] - res['n_ok']} entraxe(s) sur {res['n_total']} sans "
                              "engrènement possible : voir le tableau détaillé.").runs[0].bold = True
        ok_jeu = optimisation is None or res["jeu_min"] >= optimisation.get("jeu_min_admis", 0)
        _tableau(doc, ["Critère", "Pire cas sur la plage"], [
            ["Variation de couple maximale", f"{res['varC_max']:.2f} %"],
            ["Rendement minimal", f"{res['rendement_min']:.2f} %"],
            ["Pression de Hertz maximale (1 mN·m)", f"{res['hertz_max']:.0f} N/mm²"],
            ["Jeu de la menée", f"{res['jeu_min']:.3f}° à {res['jeu_max']:.3f}°"],
        ], largeurs_cm=[9.5, 6.0])
        doc.add_paragraph()
        doc.add_paragraph(("✓ " if ok_jeu else "⚠ ") + (
            "le jeu reste au-dessus du minimum admissible sur toute la plage."
            if ok_jeu else "le jeu descend sous le minimum admissible sur une partie de la plage."))

    # -- 4. tableau détaillé --------------------------------------------------------------
    _titre(doc, "4. Résultats détaillés (pas de 10 µm)")
    corps = []
    for l in lignes:
        if "stats" in l:
            st = l["stats"]
            corps.append([f"{l['da_um']:+.0f}", f"{st['couple'][3]:.2f}", f"{st['rendement'][1]:.2f}",
                         f"{l['jeu_min']:.2f}", f"{l['jeu_max']:.2f}", f"{l['hertz_max']:.0f}"])
        else:
            corps.append([f"{l['da_um']:+.0f}", "—", "—", "—", "—", f"({l['erreur']})"])
    _tableau(doc, ["ΔA (µm)", "Var. couple (%)", "Rendement moy. (%)", "Jeu min (°)", "Jeu max (°)",
                   "Hertz max (N/mm²)"], corps, largeurs_cm=[2.3, 2.6, 2.9, 2.3, 2.3, 3.1])

    # -- 5. graphiques ----------------------------------------------------------------------
    _titre(doc, "5. Graphiques")
    reperes = [("nominal", 0.0), ("minimal", da_neg), ("maximal", da_pos)]
    for nom, da in reperes:
        ligne = min(lignes, key=lambda l: abs(l["da_um"] - da))
        if "sim" not in ligne:
            doc.add_paragraph(f"Entraxe {nom} (ΔA = {da:+g} µm) : pas d'engrènement possible.")
            continue
        png = dossier_images / f"engrenement_{nom}.png"
        figure_engrenement(png, roues, profils, mobiles, ligne["sim"], ligne["da_mm"],
                           f"Entraxe {nom} (ΔA = {ligne['da_um']:+.0f} µm)")
        doc.add_picture(str(png), width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    png_tend = dossier_images / "tendance.png"
    figure_tendance(png_tend, lignes, optimisation.get("jeu_min_admis") if optimisation else None,
                    optimisation.get("jeu_max_admis") if optimisation else None)
    doc.add_picture(str(png_tend), width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # -- 6. géométrie (annexe) --------------------------------------------------------------
    _titre(doc, "6. Géométrie du profil (annexe)")
    for nom, m in (("Menante", m1), ("Menée", m2)):
        doc.add_paragraph(nom).runs[0].bold = True
        pm, ra = m.arcs()
        corps = [[i + 1, f"{x:.5f}", f"{y:.5f}"] for i, (x, y) in enumerate(m.points)]
        for i, r in enumerate(ra):
            corps[i].append(f"{r:+.4f}")
        corps[-1].append("-")
        _tableau(doc, ["Point", "x (mm)", "y (mm)", "Rayon de l'arc suivant (mm)"], corps,
                 largeurs_cm=[2.0, 4.0, 4.0, 5.5])
        doc.add_paragraph()

    # -- 7. propositions ----------------------------------------------------------------
    _titre(doc, "7. Remarques et propositions")
    for txt in (
        "Ce rapport ne contient pas encore de comparaison chiffrée avec la sortie du programme "
        "Fortran ENGCRE : je peux ajouter un tableau côte à côte si vous me communiquez les "
        "valeurs de référence (ou le fichier de sortie).",
        "Le contrôle d'interférence entre dents (utilisé pendant l'optimisation) n'est pas "
        "repris ici : je peux ajouter une colonne « interférence » au tableau détaillé si "
        "c'est utile pour la validation finale.",
        "Le rayon minimal de l'arc de tête (durée de vie des fraises) est documenté dans les "
        "conditions d'optimisation, mais pas mesuré directement sur ce profil ; je peux "
        "l'ajouter à l'annexe géométrie.",
        "Une comparaison avant/après optimisation (le profil développante d'origine, à côté du "
        "profil optimisé) rendrait le gain plus visible qu'une simple valeur de J.",
        "Si ce rapport doit être transmis à un tiers, un paragraphe de conclusion signé "
        "(nom, fonction, date) et un numéro de révision seraient à ajouter.",
    ):
        doc.add_paragraph(txt, style="List Bullet")

    doc.save(chemin)
    return lignes, res
