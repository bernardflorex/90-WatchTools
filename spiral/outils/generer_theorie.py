# outils/generer_theorie.py
"""
Génère « Spiral_Lab_Theorie_et_Structure.docx », le document de référence
destiné à la VALIDATION CROISÉE : il doit permettre de reconstruire le modèle
dans un autre outil de calcul et de comparer les résultats poste par poste.

Pourquoi un générateur versionné plutôt qu'un document édité à la main : la
première version (26/09/2026) avait été produite au fil d'une session, sans
script conservé. Trois jours plus tard, le modèle avait bougé -- repère du
mouvement, critère DVH, objectif d'optimisation -- et il a fallu extraire tout
le texte du .docx pour le reconstruire. Le contenu est ici de la DONNÉE (la
liste CONTENU), le rendu Word n'étant qu'une fonction : mettre le document à
jour revient à éditer du texte, et la mise en page ne peut plus diverger d'une
section à l'autre.

    python outils/generer_theorie.py [chemin_de_sortie]
"""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Largeur utile en DXA (1/1440 de pouce) : A4 moins les deux marges. Les
# tableaux y sont normalisés, faute de quoi ils débordent de la page -- Word ne
# les recadre pas tout seul quand la somme des colonnes dépasse.
MARGE_CM = 2.0
UTILE_DXA = int(round((21.0 - 2 * MARGE_CM) / 2.54 * 1440))
BLEU = RGBColor(0x1F, 0x3B, 0x5C)
GRIS = RGBColor(0x66, 0x66, 0x66)


# ----------------------------------------------------------------- rendu
def _titre(doc, texte, niveau):
    p = doc.add_heading(texte, level=niveau)
    for r in p.runs:
        r.font.color.rgb = BLEU
    return p


def _para(doc, texte, gras_avant=None):
    """Paragraphe courant. `gras_avant` met en gras une amorce du type
    « Justification mesurée. » : c'est ce qui rend le document parcourable en
    diagonale, chaque paragraphe annonçant sa nature."""
    p = doc.add_paragraph()
    if gras_avant:
        r = p.add_run(gras_avant)
        r.bold = True
    p.add_run(texte)
    return p


def _puce(doc, texte):
    return doc.add_paragraph(texte, style="List Bullet")


def _eq(doc, lignes):
    """Bloc d'équations ou de pseudo-code, en chasse fixe et indenté.

    Les formules sont écrites en texte ASCII plutôt qu'en objets équation Word :
    le document doit pouvoir être relu, copié et transcrit dans un autre outil
    sans dépendre du moteur d'équations, et les indices Unicode suffisent à la
    lisibilité."""
    for ligne in lignes:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(ligne)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def _tableau(doc, entetes, lignes, poids=None):
    """Tableau normalisé sur la largeur utile. `poids` donne la répartition
    relative des colonnes ; sans lui, elles sont égales."""
    n = len(entetes)
    poids = poids or [1] * n
    total = float(sum(poids))
    t = doc.add_table(rows=1, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (cellule, texte) in enumerate(zip(t.rows[0].cells, entetes)):
        cellule.text = ""
        r = cellule.paragraphs[0].add_run(texte)
        r.bold = True
        r.font.size = Pt(9)
    for ligne in lignes:
        cells = t.add_row().cells
        for cellule, texte in zip(cells, ligne):
            cellule.text = ""
            r = cellule.paragraphs[0].add_run(str(texte))
            r.font.size = Pt(9)
    # Largeurs posées sur la grille XML : régler cell.width seul ne suffit pas,
    # Word repartant de la grille pour recalculer la mise en page.
    grille = t._tbl.find(qn("w:tblGrid"))
    for col, p in zip(grille.findall(qn("w:gridCol")), poids):
        col.set(qn("w:w"), str(int(round(UTILE_DXA * p / total))))
    for rang in t.rows:
        for cellule, p in zip(rang.cells, poids):
            cellule.width = Cm((21.0 - 2 * MARGE_CM) * p / total)
    return t


RENDU = {
    "h1": lambda doc, a: _titre(doc, a[0], 1),
    "h2": lambda doc, a: _titre(doc, a[0], 2),
    "p": lambda doc, a: _para(doc, a[0], a[1] if len(a) > 1 else None),
    "puce": lambda doc, a: _puce(doc, a[0]),
    "eq": lambda doc, a: _eq(doc, a[0]),
    "tab": lambda doc, a: _tableau(doc, a[0], a[1], a[2] if len(a) > 2 else None),
    "saut": lambda doc, a: doc.add_page_break(),
}


def construire(contenu, chemin):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(MARGE_CM)
    s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)

    titre = doc.add_heading(
        "Spiral Lab — théorie, outils mathématiques et structure du programme", level=0)
    for r in titre.runs:
        r.font.color.rgb = BLEU
    for texte, taille, couleur in (
        ("Optimisation de la géométrie d'un spiral plat en silicium", 12, BLEU),
        ("Modèle physique, méthodes numériques et justification des choix", 11, GRIS),
        ("Document de référence pour la validation croisée — 30 septembre 2026", 10, GRIS),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texte)
        r.font.size = Pt(taille)
        r.font.color.rgb = couleur

    for bloc in contenu:
        RENDU[bloc[0]](doc, bloc[1:])

    doc.save(chemin)
    return doc


if __name__ == "__main__":
    from contenu_theorie import CONTENU

    chemin = sys.argv[1] if len(sys.argv) > 1 else "Spiral_Lab_Theorie_et_Structure.docx"
    d = construire(CONTENU, chemin)
    n_h1 = sum(1 for p in d.paragraphs if p.style.name == "Heading 1")
    n_h2 = sum(1 for p in d.paragraphs if p.style.name == "Heading 2")
    print(f"{chemin} : {n_h1} sections, {n_h2} sous-sections, "
          f"{len(d.tables)} tableaux, {len(d.paragraphs)} paragraphes")
