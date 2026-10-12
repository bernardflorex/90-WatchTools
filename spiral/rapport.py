# rapport.py
"""
Enregistrement d'un projet : un dossier contenant tout ce qui décrit une
simulation — le rapport, l'animation, les fichiers de gravure et la
configuration.

Le rapport est un document Word. L'animation y figure, mais Word n'anime pas
les GIF : il n'en montre que la première image. C'est pourquoi le GIF est aussi
déposé comme fichier autonome à côté du rapport, où il s'anime normalement ;
le rapport renvoie vers lui nommément.

Le tout est écrit dans un sous-dossier de `Results/` portant le nom du projet,
de sorte qu'une simulation se transmette ou s'archive d'un bloc, sans avoir à
rassembler des fichiers dispersés.
"""
import io
import os
from datetime import datetime

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from matplotlib.figure import Figure

from commun import theme
from . import dxf_io
from .constantes import RHO_SILICIUM
from .echappement import ReglageEchappement
from .isochronisme import (POSITIONS_HORIZONTALES, analyser_marches,
                            positions_verticales, reperes_horaires)
from .physique import (angle_point_comptage, axe_force_pivot, frequence_propre,
                        marche_raquette, resoudre_elastica)
from .presets import sauvegarder_configuration
from .rendu import (calculer_frames_deploiement, sauvegarder_gif_deploiement,
                     tracer_repere_mouvement, tracer_vecteur_g)
from .thermique import ReglageThermique, analyser_thermique, etude_sensibilite

# Largeur utile en DXA (1/1440 de pouce) : A4 moins les deux marges. Les
# tableaux y sont normalisés, faute de quoi ils débordent de la page.
_MARGE_CM = 2.0
_UTILE_DXA = int(round((21.0 - 2 * _MARGE_CM) / 2.54 * 1440))
_UTILE_CM = 21.0 - 2 * _MARGE_CM


# ----------------------------------------------------------------- figures
def _fig_etat(geom, r_min, angle_deg, titre, limites, jeu_um=None,
               reperes_positions=False, cote_rouage=True, angle_collet_deg=None,
               angle_ancre_deg=None):
    """Un état du spiral, la forme au repos rappelée en gris.

    Toutes les vues partagent les MÊMES limites d'axes : sans cela l'oeil
    interprète un changement d'échelle comme un déplacement du centre de
    rotation (même piège que sur l'animation)."""
    fig = Figure(figsize=(4.8, 4.8))
    ax = fig.add_subplot(111)
    ext0, int0 = geom.contour_flancs()
    ax.plot(ext0[:, 0], ext0[:, 1], "-", color="0.82", lw=0.7)
    ax.plot(int0[:, 0], int0[:, 1], "-", color="0.82", lw=0.7)
    if angle_deg is None:
        # Vue au repos : la forme non déformée EST le sujet, elle est donc
        # retracée par-dessus dans la couleur des autres vues. Sans cela la
        # vue n'aurait que le gris pâle de référence et paraîtrait délavée.
        ax.plot(ext0[:, 0], ext0[:, 1], "-", color="tab:blue", lw=0.9)
        ax.plot(int0[:, 0], int0[:, 1], "-", color="tab:blue", lw=0.9)
    else:
        Xe, Ye, Xi, Yi, _, _ = resoudre_elastica(geom, float(angle_deg), r_min)
        ax.plot(Xe, Ye, "-", color="tab:blue", lw=0.9)
        ax.plot(Xi, Yi, "-", color="tab:blue", lw=0.9)
    ax.plot(0, 0, "k+", ms=10)

    if reperes_positions:
        tracer_repere_mouvement(ax, float(geom.rayon().max()), cote_rouage,
                                 angle_collet_deg=angle_collet_deg,
                                 rayon_min=float(geom.rayon().min()),
                                 angle_ancre_deg=angle_ancre_deg)

    if reperes_positions:
        # Les repères horaires sont posés à 1.10 rayon : la vue qui les porte a
        # besoin d'un peu plus de champ que les trois autres, sans quoi les
        # étiquettes sortent du cadre. L'échelle reste commune (même centre,
        # même unité par pixel) -- seule la marge change.
        r_ = 1.20 * float(geom.rayon().max())
        cx = 0.5 * (limites[0] + limites[1])
        cy = 0.5 * (limites[2] + limites[3])
        ax.set_xlim(min(limites[0], cx - r_), max(limites[1], cx + r_))
        ax.set_ylim(min(limites[2], cy - r_), max(limites[3], cy + r_))
    else:
        ax.set_xlim(limites[0], limites[1])
        ax.set_ylim(limites[2], limites[3])
    ax.set_aspect("equal")
    ax.set_xlabel("x (mm)")
    ax.set_ylabel("y (mm)")
    ax.set_title(titre + (f"  —  jeu {jeu_um:.1f} µm" if jeu_um is not None else ""),
                  fontsize=10)
    fig.tight_layout()
    return fig


def _fig_position_cosc(geom, r_min, nom_position, cote_rouage=True,
                        angle_ancre_deg=None):
    """Le spiral tel qu'il est PHYSIQUEMENT ORIENTÉ dans une position COSC.

    Le dessin est tourné pour que la gravité aille vers le bas de la page : la
    pesanteur ne tourne pas, c'est la montre qu'on incline. La position COSC
    « 6H » est celle où le repère 6H est EN HAUT -- donc où g pointe vers le
    12H -- et non l'inverse.
    """
    angle_g = positions_verticales(cote_rouage)[nom_position]
    # On amène g sur 270° (le bas de la page) : tout le mouvement pivote de
    # cet angle, repères horaires et ligne d'échappement compris.
    rot = np.radians(270.0) - angle_g
    c, sn = np.cos(rot), np.sin(rot)

    def tourner(x, y):
        return x * c - y * sn, x * sn + y * c

    fig = Figure(figsize=(4.8, 4.8))
    ax = fig.add_subplot(111)
    ext0, int0 = geom.contour_flancs()
    for contour in (ext0, int0):
        xr, yr = tourner(contour[:, 0], contour[:, 1])
        ax.plot(xr, yr, "-", color="tab:blue", lw=0.9)
    ax.plot(0, 0, "k+", ms=10)

    rayon = float(geom.rayon().max())
    for nom, a in reperes_horaires(cote_rouage).items():
        ax.annotate(nom, (1.16 * rayon * np.cos(a + rot), 1.16 * rayon * np.sin(a + rot)),
                    color="0.35", fontsize=8.5, ha="center", va="center")
    a_tige = reperes_horaires(cote_rouage)["3H"] + rot
    ax.plot([0.0, 1.00 * rayon * np.cos(a_tige)], [0.0, 1.00 * rayon * np.sin(a_tige)],
            "-", color="0.55", lw=1.2, zorder=1)
    if angle_ancre_deg is not None:
        # Demi-axe seulement, vers l'ancre : il tourne avec le mouvement, comme
        # tout le reste du dessin.
        a = np.radians(angle_ancre_deg) + rot
        ax.plot([0.0, 1.06 * rayon * np.cos(a)], [0.0, 1.06 * rayon * np.sin(a)],
                "--", color="tab:green", lw=1.2, zorder=1)

    tracer_vecteur_g(ax, rayon, nom_position)
    lim = 1.55 * rayon
    ax.set_xlim(-lim, lim)
    ax.set_ylim(-lim, lim)
    ax.set_aspect("equal")
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_title(f"Position COSC {nom_position}", fontsize=10)
    fig.tight_layout()
    return fig


def _limites_communes(geom, r_min, angles):
    xs, ys = [], []
    ext0, int0 = geom.contour_flancs()
    xs += [ext0[:, 0].min(), ext0[:, 0].max()]
    ys += [ext0[:, 1].min(), ext0[:, 1].max()]
    for a in angles:
        Xe, Ye, Xi, Yi, _, _ = resoudre_elastica(geom, float(a), r_min)
        xs += [min(Xe.min(), Xi.min()), max(Xe.max(), Xi.max())]
        ys += [min(Ye.min(), Yi.min()), max(Ye.max(), Yi.max())]
    x0, x1, y0, y1 = min(xs), max(xs), min(ys), max(ys)
    # Marge généreuse : elle doit laisser la place aux étiquettes de position
    # et à celle du vecteur g, qui débordent du spiral.
    marge = 0.20 * max(x1 - x0, y1 - y0)
    return x0 - marge, x1 + marge, y0 - marge, y1 + marge


def _fig_orbite(diag):
    fig = Figure(figsize=(8.6, 6.6))
    gs = fig.add_gridspec(2, 2, height_ratios=[1, 1.4], width_ratios=[40, 1])
    ax = fig.add_subplot(gs[0, 0])
    cax = fig.add_subplot(gs[0, 1])
    ax2 = fig.add_subplot(gs[1, :])
    ang = diag["angles_deg"]
    xc, yc = diag["x_cg_mm"] * 1000.0, diag["y_cg_mm"] * 1000.0
    tr = ax.scatter(xc, yc, c=ang, cmap="coolwarm", s=14, zorder=3)
    ax.plot(xc, yc, "-", color="0.6", lw=0.8, zorder=2)
    ax.plot(xc.mean(), yc.mean(), "k+", ms=12, zorder=4)
    ax.set_aspect("equal", adjustable="datalim")
    demi = max(float(np.ptp(xc)), float(np.ptp(yc)), 1e-6) * 0.6
    ax.set_xlim(0.5 * (xc.max() + xc.min()) - demi, 0.5 * (xc.max() + xc.min()) + demi)
    ax.set_ylim(0.5 * (yc.max() + yc.min()) - demi, 0.5 * (yc.max() + yc.min()) + demi)
    ax.set_xlabel("x du CG (µm)")
    ax.set_ylabel("y du CG (µm)")
    ax.grid(True, linestyle=":", alpha=0.5)
    rayon = float(np.hypot(xc - xc.mean(), yc - yc.mean()).max())
    ax.set_title(f"Orbite du centre de gravité — rayon max {rayon:.1f} µm, "
                 f"décentrage moyen {np.hypot(xc.mean(), yc.mean()):.1f} µm", fontsize=10)
    barre = fig.colorbar(tr, cax=cax)
    barre.set_label("angle du balancier (°)", fontsize=8)
    ax2.plot(ang, xc - xc.mean(), "o-", ms=3, label="x (écart à la moyenne)")
    ax2.plot(ang, yc - yc.mean(), "s-", ms=3, label="y (écart à la moyenne)")
    ax2.axvline(0, color="k", lw=0.8)
    ax2.set_xlabel("Angle du balancier (°)")
    ax2.set_ylabel("Déplacement du CG (µm)")
    ax2.grid(True, linestyle=":", alpha=0.5)
    ax2.legend(fontsize=8)
    ax2.set_title(f"{len(ang)} positions de déploiement calculées", fontsize=10)
    fig.tight_layout()
    return fig


def _fig_marche(resultat, reglage):
    fig = Figure(figsize=(8.6, 4.8))
    ax = fig.add_subplot(111)
    amp = resultat["amplitudes_deg"]
    for nom, val in resultat["totales"].items():
        horizontale = nom in POSITIONS_HORIZONTALES
        ax.plot(amp, val, "s--" if horizontale else "o-",
                lw=2.0 if horizontale else 1.0, ms=4, label=nom)
    ax.plot(amp, resultat["isochronisme"], ":", color="0.35", lw=1.5, label="isochronisme propre")
    ax.plot(amp, resultat["echappement"], ":", color="tab:red", lw=1.5, label="échappement")
    if resultat["diag_jeu_goupilles"].get("actif"):
        ax.plot(amp, resultat["jeu_goupilles"], ":", color="tab:purple", lw=1.5,
                label="jeu de goupilles")
    ax.axvline(reglage.amplitude_basse_deg, color="0.7", ls="-.", lw=0.8)
    ax.axhline(0, color="k", lw=0.8)
    ax.set_xlabel("Amplitude (deg)")
    ax.set_ylabel("Marche diurne (s/j)")
    ax.grid(True, linestyle=":", alpha=0.5)
    ax.legend(fontsize=8, ncol=3)
    fig.tight_layout()
    return fig


def _fig_epaisseur(geom):
    fig = Figure(figsize=(8.6, 3.2))
    ax = fig.add_subplot(111)
    if geom.epaisseur_oxyde > 0:
        ax.plot(geom.s, geom.epaisseur_finie * 1000, "-", label="lame finie (oxyde compris)")
        ax.plot(geom.s, geom.epaisseur * 1000, "--", color="0.5", label="cote de gravure")
        ax.legend(fontsize=8)
    else:
        ax.plot(geom.s, geom.epaisseur * 1000, "-")
    if geom.a_partie_morte:
        ax.axvline(geom.s[geom.n_actif - 1], color="tab:red", ls="-.", lw=1.0)
    ax.set_xlabel("Longueur depuis la virole (mm)")
    ax.set_ylabel("Épaisseur (µm)")
    ax.grid(True, linestyle=":", alpha=0.6)
    fig.tight_layout()
    return fig


def _fig_thermique(geom, inertie, reglage_th):
    a_nu = analyser_thermique(geom, inertie, reglage_th, t_oxyde_um=0.0)
    t = geom.epaisseur_oxyde * 1000.0
    a_ox = analyser_thermique(geom, inertie, reglage_th, t_oxyde_um=t if t > 0 else None)
    fig = Figure(figsize=(8.6, 4.0))
    ax = fig.add_subplot(111)
    ax.plot(a_nu["temperatures"], a_nu["marches"], "o-", color="tab:red",
            label=f"silicium nu ({a_nu['pente_s_par_j_par_degc']:+.2f} s/j/°C)")
    if a_ox["valide"]:
        ax.plot(a_ox["temperatures"], a_ox["marches"], "s-", color="tab:blue",
                label=f"oxydé {a_ox['epaisseur_oxyde_um']:.2f} µm "
                      f"({a_ox['pente_s_par_j_par_degc']:+.3f} s/j/°C)")
    etendue = reglage_th.temperature_max - reglage_th.temperature_min
    ax.axhspan(-0.6 * etendue / 2, 0.6 * etendue / 2, color="0.9", zorder=0,
               label="bande COSC (0.6 s/j/°C)")
    ax.axhline(0, color="k", lw=0.8)
    ax.set_xlabel("Température (°C)")
    ax.set_ylabel("Marche (s/j)")
    ax.grid(True, linestyle=":", alpha=0.5)
    ax.legend(fontsize=8)
    fig.tight_layout()
    return fig, a_nu, a_ox


# ----------------------------------------------------------------- document
def _inserer_figure(doc, fig, largeur_cm=None, legende=None):
    tampon = io.BytesIO()
    fig.savefig(tampon, format="png", dpi=110, bbox_inches="tight")
    tampon.seek(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(tampon, width=Cm(largeur_cm or _UTILE_CM))
    if legende:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(legende)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _grille_dxa(t, largeurs):
    """Fixe tblW et tblGrid : Word se fie à la grille, pas seulement aux
    largeurs de cellules."""
    tbl = t._tbl
    tblPr = tbl.tblPr
    for ancien in tblPr.findall(qn("w:tblW")):
        tblPr.remove(ancien)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(sum(largeurs)))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    for ancienne in tbl.findall(qn("w:tblGrid")):
        tbl.remove(ancienne)
    grille = OxmlElement("w:tblGrid")
    for lg in largeurs:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(lg))
        grille.append(col)
    tblPr.addnext(grille)


def _tableau(doc, entetes, lignes, proportions):
    """Les proportions sont normalisées sur la largeur utile : un tableau
    dimensionné en absolu déborde de la page."""
    total = float(sum(proportions))
    largeurs = [int(round(x / total * _UTILE_DXA)) for x in proportions]
    largeurs[-1] += _UTILE_DXA - sum(largeurs)
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _grille_dxa(t, largeurs)
    for i, e in enumerate(entetes):
        cel = t.rows[0].cells[i]
        cel.text = ""
        r = cel.paragraphs[0].add_run(str(e))
        r.bold = True
        r.font.size = Pt(9.5)
    for ligne in lignes:
        cells = t.add_row().cells
        for i, v in enumerate(ligne):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(v))
            r.font.size = Pt(9)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Emu(largeurs[i] * 635)
    return t


def _encadre(doc, titre, texte, couleur):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run(f"{titre} ")
    r.bold = True
    r.font.color.rgb = couleur
    r2 = p.add_run(texte)
    r2.font.color.rgb = couleur


@theme.clair()
def construire_rapport(chemin_docx: str, geom, spec, parametres,
                        reglage_echappement: ReglageEchappement = None,
                        reglage_thermique: ReglageThermique = None,
                        amplitudes_deg=None, pas_balayage_deg: float = 20.0,
                        jeu_goupilles_um: float = 0.0, r_min: float = None,
                        chemin_gif: str = None, origine: str = "",
                        progression=None) -> str:
    """Écrit le rapport Word et retourne son chemin.

    `chemin_gif` : animation à insérer. Word n'anime pas les GIF -- seule la
    première image apparaît -- donc le rapport renvoie explicitement vers le
    fichier autonome déposé à côté, qui lui s'anime."""
    def etape(texte):
        if progression is not None:
            progression(texte)

    reglage_echappement = reglage_echappement or ReglageEchappement()
    reglage_thermique = reglage_thermique or ReglageThermique()
    if amplitudes_deg is None:
        amplitudes_deg = np.arange(200.0, spec.amplitude_max + 1e-6, 20.0)
    if r_min is None:
        r_min = float(geom.rayon()[0])
    cote_rouage = getattr(spec, "vue_cote_rouage", True)
    # Même règle que ParametresSpiral.angle_collet_deg : l'angle de la ligne des
    # centres ne la définit qu'à 180° près, `collet_cote_ancre` tranche lequel
    # de ses deux bouts porte le collet. `getattr` pour rester tolérant à une
    # spec sauvegardée avant l'introduction du repère du mouvement.
    angle_ancre_deg = ((getattr(spec, "angle_echappement_deg", -6.2) + 180.0)
                        % 360.0 - 180.0)
    angle_collet_deg = angle_ancre_deg
    if not getattr(spec, "collet_cote_ancre", True):
        angle_collet_deg += 180.0
    angle_collet_deg = (angle_collet_deg + 180.0) % 360.0 - 180.0

    etape("Balayage elastica et marche diurne...")
    resultat = analyser_marches(geom, r_min, spec.freq_cible, spec.inertie_balancier,
                                amplitudes_deg, reglage_echappement=reglage_echappement,
                                positions=positions_verticales(cote_rouage),
                                pas_balayage_deg=pas_balayage_deg,
                                jeu_goupilles_um=jeu_goupilles_um)
    diag = resultat["diag_orbite"]
    angle_jeu_min = diag["angle_jeu_min_deg"]
    jeu_min_um = diag["jeu_min_mm"] * 1000.0
    freq = frequence_propre(geom, spec.inertie_balancier)
    module_force, axe_force = axe_force_pivot(geom, r_min, spec.amplitude_max)

    etape("États caractéristiques du spiral...")
    # Fermeture = rotation NÉGATIVE : le rayon moyen diminue, les spires se
    # resserrent (vérifié : 1.317 mm à -340° contre 1.572 mm à +340°).
    fermeture, ouverture = -spec.amplitude_max, +spec.amplitude_max
    limites = _limites_communes(geom, r_min, [fermeture, ouverture, angle_jeu_min])
    etats = [
        ("Au repos (balancier à zéro)", None, None, True),
        (f"Fermeture complète ({fermeture:+.0f}°)", fermeture, None, False),
        (f"Ouverture complète ({ouverture:+.0f}°)", ouverture, None, False),
        (f"Jeu inter-spires minimal ({angle_jeu_min:+.0f}°)", angle_jeu_min, jeu_min_um, False),
    ]

    etape("Orbite, marche, épaisseur, température...")
    fig_orbite = _fig_orbite(diag)
    fig_marche = _fig_marche(resultat, reglage_echappement)
    fig_epaisseur = _fig_epaisseur(geom)
    fig_therm, a_nu, a_ox = _fig_thermique(geom, spec.inertie_balancier, reglage_thermique)
    sensibilite = etude_sensibilite(geom, spec.inertie_balancier, reglage_thermique)

    etape("Rédaction du document...")
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(_MARGE_CM)
    s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)

    def h(txt, niveau=1):
        p = doc.add_heading(txt, level=niveau)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x5C)

    titre = doc.add_heading("Rapport de simulation — spiral horloger", level=0)
    for r in titre.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x5C)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{datetime.now():%d/%m/%Y à %H:%M}"
                  + (f" — géométrie : {origine}" if origine else ""))
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if jeu_min_um < spec.jeu_min_spires * 1000.0:
        _encadre(doc, "Jeu inter-spires insuffisant :",
                 f"{jeu_min_um:.1f} µm mesuré pour {spec.jeu_min_spires*1000:.1f} µm requis. "
                 f"Les spires se touchent à {angle_jeu_min:+.0f}°.", RGBColor(0xC0, 0x00, 0x00))
    if abs(resultat["ecart_compensation"]) > reglage_echappement.tolerance_s_par_j:
        _encadre(doc, "Compensation hors tolérance :",
                 f"écart de {resultat['ecart_compensation']:+.2f} s/j pour une tolérance de "
                 f"±{reglage_echappement.tolerance_s_par_j:.2f} s/j.", RGBColor(0xA0, 0x6A, 0x00))

    # ---------------- 1. configuration ----------------
    h("1. Configuration", 1)
    h("Mouvement", 2)
    _tableau(doc, ["Paramètre", "Valeur"], [
        ("Rayon au collet r_min", f"{spec.r_min:.4f} mm"),
        ("Rayon au piton r_max", f"{spec.r_max:.4f} mm"),
        ("Nombre de tours réels (collet → point de comptage)", f"{spec.n_tours_reels:g}"),
        ("Nombre de tours fictifs (pas de référence)", f"{spec.n_tours_fictifs:g}"),
        ("Hauteur hors-plan", f"{spec.hauteur_h*1000:.1f} µm"),
        ("Fréquence cible", f"{spec.freq_cible:.4f} Hz"),
        ("Inertie du balancier", f"{spec.inertie_balancier:.4g} kg·m²"),
        ("Jeu inter-spires minimal requis", f"{spec.jeu_min_spires*1000:.1f} µm"),
        ("Amplitude maximale", f"{spec.amplitude_max:.0f}°"),
    ], [60, 40])

    p_ = parametres
    lignes_forme = [
        ("Éléments de maillage", f"{geom.n}"),
        ("Grossmann — angle / n_pas / n_points",
         f"{p_.grossmann_angle_deg:.0f}° / {p_.grossmann_n_pas} / {p_.grossmann_n_points}"),
    ]
    if p_.raquette_active:
        lignes_forme += [
            ("Terminaison extérieure", "raquette (courbe en S + arc concentrique)"),
            ("Transition — angle / n_points",
             f"{p_.raquette_transition_angle_deg:.0f}° / {p_.raquette_n_points}"),
            ("Arc avant le point de comptage", f"{p_.raquette_arc_avant_deg:.0f}°"),
            ("Arc du point de comptage au piton", f"{p_.raquette_piton_angle_deg:.0f}°"),
            ("Plage de réglage exigée", f"±{p_.raquette_plage_s_par_j:.0f} s/j"),
            ("Jeu libre entre goupilles", f"{p_.raquette_jeu_um:.1f} µm"),
        ]
    else:
        lignes_forme += [
            ("Terminaison extérieure", "spiral libre (courbe de Phillips)"),
            ("Phillips — angle / n_pas / n_points",
             f"{p_.phillips_angle_deg:.0f}° / {p_.phillips_n_pas} / {p_.phillips_n_points}"),
        ]
    lignes_forme += [
        ("Profil d'épaisseur — nœuds", f"{len(p_.epaisseur_noeuds)}"),
        ("Écart radial — nœuds", f"{len(p_.delta_r_noeuds)}"),
        ("Oxydation", "active" if p_.oxydation_active else "inactive"),
    ]
    h("Forme du spiral", 2)
    _tableau(doc, ["Paramètre", "Valeur"], lignes_forme, [60, 40])

    h("Réglage et environnement", 2)
    _tableau(doc, ["Paramètre", "Valeur"], [
        ("Retard d'échappement", f"{reglage_echappement.retard_s_par_j:.1f} s/j à "
                                 f"{reglage_echappement.amplitude_basse_deg:.0f}°"),
        ("Amplitude de référence", f"{reglage_echappement.amplitude_reference_deg:.0f}°"),
        ("Tolérance de compensation", f"±{reglage_echappement.tolerance_s_par_j:.2f} s/j"),
        ("Montage visé", reglage_echappement.variante_visee),
        ("Plage de température",
         f"{reglage_thermique.temperature_min:.0f} à {reglage_thermique.temperature_max:.0f} °C "
         f"(référence {reglage_thermique.temperature_reference:.0f} °C)"),
        ("Dilatation du balancier", f"{reglage_thermique.cte_balancier*1e6:.1f} ppm/K"),
        ("Pas de balayage angulaire", f"{pas_balayage_deg:.0f}°"),
    ], [60, 40])

    # ---------------- 2. résultats ----------------
    doc.add_page_break()
    h("2. Résultats numériques", 1)
    resultats = [
        ("Fréquence propre", f"{freq:.5f} Hz",
         f"écart à la cible {86400*(freq/spec.freq_cible-1):+.1f} s/j"),
        ("Masse du spiral", f"{geom.masse_totale(RHO_SILICIUM)*1e9:.2f} µg", ""),
        ("Longueur développée", f"{geom.longueur:.2f} mm",
         f"active {geom.partie_active().longueur:.2f} mm"),
        ("Épaisseur de lame",
         f"{geom.epaisseur_finie.min()*1000:.2f} à {geom.epaisseur_finie.max()*1000:.2f} µm",
         "cote finie, oxyde compris" if geom.epaisseur_oxyde > 0 else ""),
        ("Jeu inter-spires minimal", f"{jeu_min_um:.1f} µm",
         f"à {angle_jeu_min:+.0f}° — requis {spec.jeu_min_spires*1000:.1f} µm"),
        ("Dispersion de marche (gravité, verticales)",
         f"{resultat['dispersion_gravite']:.2f} s/j", "écart max entre positions verticales"),
        ("DVH (critère COSC)", f"{resultat['dvh_reference']:+.2f} s/j",
         f"position 6H moins horizontale, à {resultat['amplitude_dvh_deg']:.0f}° — "
         f"pire cas sur la plage {resultat['dvh_max']:.2f} s/j"),
        ("Vecteur de sensibilité à la gravité",
         f"{resultat['sensibilite_module_reference']:.2f} s/j",
         f"phase {resultat['sensibilite_phase_deg']:.1f}° — "
         f"marche(θ) = A·cos(θ − φ) sur toutes les verticales"),
        (f"Isochronisme propre à {reglage_echappement.amplitude_basse_deg:.0f}°",
         f"{resultat['iso_amplitude_basse']:+.2f} s/j",
         f"cible {reglage_echappement.retard_s_par_j:+.2f} s/j"),
        ("Écart de compensation", f"{resultat['ecart_compensation']:+.2f} s/j",
         "dans la tolérance" if abs(resultat["ecart_compensation"])
         <= reglage_echappement.tolerance_s_par_j else "HORS tolérance"),
        ("Force sur le pivot du balancier", f"{module_force:.1f} µN",
         f"axe {axe_force:.1f}° — point de comptage {angle_point_comptage(geom):.1f}°"),
    ]
    if geom.a_partie_morte:
        rap, lent, etendue = marche_raquette(geom, spec.inertie_balancier)
        contact = resultat["diag_jeu_goupilles"].get("amplitude_contact_deg", float("nan"))
        resultats += [
            ("Plage de la raquette", f"{rap:+.1f} / {lent:+.1f} s/j", f"étendue {etendue:.1f} s/j"),
            ("Retard dû au jeu de goupilles", f"{resultat['jeu_amplitude_basse']:+.2f} s/j",
             f"contact dès {contact:.0f}° de balancier"),
        ]
    if geom.epaisseur_oxyde > 0:
        resultats.append(("Épaisseur d'oxyde", f"{geom.epaisseur_oxyde*1000:.3f} µm",
                          f"lame élargie de {a_ox['croissance_lame_um']:.2f} µm à la gravure"))
    resultats.append(("Marche thermique après compensation",
                      f"{a_ox['pente_s_par_j_par_degc']:+.4f} s/j/°C" if a_ox["valide"] else "—",
                      f"silicium nu : {a_nu['pente_s_par_j_par_degc']:+.3f} s/j/°C"))
    _tableau(doc, ["Grandeur", "Valeur", "Remarque"], resultats, [42, 22, 36])

    # ---------------- 3. états ----------------
    doc.add_page_break()
    h("3. États caractéristiques du spiral", 1)
    doc.add_paragraph(
        "La forme au repos est rappelée en gris clair sur chaque vue. Les quatre vues "
        "partagent les mêmes limites d'axes : un changement d'échelle donnerait "
        "l'impression que le centre de rotation se déplace. Elles sont tracées dans le "
        "repère du MOUVEMENT vu côté rouage, 6H en haut et tige à droite, et la première "
        "porte les repères horaires, "
        "l'axe de tige, la ligne droite de l'échappement et le côté depuis lequel la "
        "pièce est vue. Ce dernier engage une chiralité : se tromper de côté échange 3H "
        "et 9H sans rien changer à la dispersion, mais attribue les valeurs aux "
        "mauvaises positions dès qu'on les compare à un relevé.")
    for t_, angle, jeu_, reperes in etats:
        _inserer_figure(doc, _fig_etat(geom, r_min, angle, t_, limites, jeu_,
                                        reperes_positions=reperes,
                                        cote_rouage=cote_rouage,
                                        angle_collet_deg=angle_collet_deg,
                                        angle_ancre_deg=angle_ancre_deg),
                        largeur_cm=11.0, legende=t_)

    h("Attitudes COSC", 2)
    doc.add_paragraph(
        "Les vues ci-dessus montrent le mouvement à plat, 6H en haut. Les quatre vues "
        "suivantes montrent la même pièce dans l'ATTITUDE de chaque position verticale "
        "COSC : le dessin est tourné pour que la pesanteur aille vers le bas, puisque "
        "c'est la montre qu'on incline et non la gravité qui change de sens. Le nom de "
        "la position est celui du repère qui se trouve EN HAUT — en position 6H, g "
        "pointe donc vers le 12H.")
    for nom in ("6H", "9H", "12H", "3H"):
        _inserer_figure(doc, _fig_position_cosc(geom, r_min, nom, cote_rouage,
                                                 angle_ancre_deg),
                        largeur_cm=8.5, legende=f"Position COSC {nom}")

    # ---------------- 4. centre de gravité ----------------
    doc.add_page_break()
    h("4. Déplacement du centre de gravité", 1)
    doc.add_paragraph(
        "C'est la grandeur que l'optimisation minimise : en position verticale, le poids "
        "du spiral exerce sur l'axe un couple proportionnel au déplacement du centre de "
        "gravité, et c'est sa variation au cours de l'oscillation — non sa position "
        "moyenne — qui fait l'erreur de marche.")
    _inserer_figure(doc, fig_orbite)

    # ---------------- 5. marche ----------------
    doc.add_page_break()
    h("5. Marche diurne", 1)
    doc.add_paragraph(
        "Les six courbes pleines ou tiretées sont la marche TOTALE par position. Les "
        "courbes en pointillés en sont les composantes indépendantes de la position : "
        "l'isochronisme propre du spiral, l'erreur d'échappement et, le cas échéant, le "
        "jeu de goupilles. Les deux positions à plat (CH et CB) sont exactement la somme "
        "des pointillés, la gravité y étant sans effet ; l'éventail entre les quatre "
        "positions verticales est le terme de gravité.")
    _inserer_figure(doc, fig_marche)
    noms = list(resultat["totales"].keys())
    _tableau(doc, ["Amplitude (°)"] + noms,
             [[f"{a:.0f}"] + [f"{resultat['totales'][n][i]:+.2f}" for n in noms]
              for i, a in enumerate(resultat["amplitudes_deg"])],
             [14] + [12] * len(noms))

    # ---------------- 6. épaisseur ----------------
    h("6. Profil d'épaisseur", 1)
    _inserer_figure(doc, fig_epaisseur,
                    legende="Depuis la virole vers le piton"
                    + (". Le trait rouge marque le point de comptage."
                       if geom.a_partie_morte else "."))

    # ---------------- 7. température ----------------
    doc.add_page_break()
    h("7. Tenue en température", 1)
    _inserer_figure(doc, fig_therm)
    h("Sensibilité aux coefficients matériau", 2)
    doc.add_paragraph(
        "Ces coefficients sont des valeurs de littérature, non des mesures. Chacun est "
        "perturbé seul dans son incertitude ; la dernière colonne donne la pente qui "
        "subsisterait si l'épaisseur d'oxyde n'était PAS recalée sur le procédé réel.")
    _tableau(doc, ["Coefficient", "Incertitude", "Épaisseur requise (µm)",
                   "Pente si non recalée (s/j/°C)"],
             [[l["coefficient"], f"±{l['incertitude']:.3g}",
               f"{l['t_moins_um']:.3f} – {l['t_plus_um']:.3f}",
               f"{l['pente_non_recalee_moins']:+.3f} … {l['pente_non_recalee_plus']:+.3f}"]
              for l in sensibilite], [28, 16, 26, 30])
    quadrature = float(np.hypot.reduce(
        [max(abs(l["pente_non_recalee_moins"]), abs(l["pente_non_recalee_plus"]))
         for l in sensibilite]))
    _encadre(doc, "Cumulées en quadrature :",
             f"{quadrature:.3f} s/j/°C si l'épaisseur calculée est appliquée sans calage "
             f"expérimental. Le calcul donne la cible, pas la certitude.",
             RGBColor(0xA0, 0x6A, 0x00))

    # ---------------- 8. animation ----------------
    if chemin_gif and os.path.exists(chemin_gif):
        doc.add_page_break()
        h("8. Animation du déploiement", 1)
        doc.add_paragraph(
            f"Va-et-vient complet entre ±{spec.amplitude_max:.0f}°. Word n'anime pas les "
            f"GIF : seule la première image apparaît ci-dessous. Le fichier animé est "
            f"déposé à côté de ce rapport, sous le nom « {os.path.basename(chemin_gif)} » ; "
            f"l'ouvrir dans un navigateur ou une visionneuse d'images pour le voir jouer.")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(chemin_gif, width=Cm(11.0))

    doc.add_paragraph()
    pied = doc.add_paragraph()
    r = pied.add_run(
        "Rapport produit par Spiral Lab. Les valeurs de marche sont calculées par "
        "résolution elastica non linéaire à chaque angle du balayage ; se reporter au "
        "document de théorie pour les hypothèses et le domaine de validité.")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    doc.save(chemin_docx)
    etape("Rapport écrit.")
    return chemin_docx


# ----------------------------------------------------------------- projet
def enregistrer_projet(racine_results: str, nom: str, geom, spec, parametres,
                        bornes=None, hyperparametres=None,
                        reglage_echappement: ReglageEchappement = None,
                        reglage_thermique: ReglageThermique = None,
                        amplitudes_deg=None, pas_balayage_deg: float = 20.0,
                        jeu_goupilles_um: float = 0.0, r_min: float = None,
                        frames_animation=None, n_frames_gif: int = 31,
                        origine: str = "", ecraser: bool = False,
                        n_points_dxf: int = None, progression=None) -> dict:
    """Écrit tout le projet dans `racine_results/nom/` : animation, rapport,
    fichiers de gravure et configuration. Retourne un dict des chemins écrits.

    Un seul dossier par simulation, pour qu'elle se transmette ou s'archive
    d'un bloc plutôt qu'en fichiers dispersés.

    Les DXF suivent l'état de la lame : DEUX fichiers si elle est oxydée
    (cote de gravure pour la DRIE, pièce finie pour le contrôle, qui diffèrent
    de ~4 µm), un seul sinon.

    n_points_dxf : densité de points des DXF, indépendante du maillage de
    calcul. Celui-ci est un compromis entre précision et temps et densifie les
    zones terminales ; un fichier de gravure n'a aucune raison d'hériter de ce
    découpage. Par défaut, le maillage est repris tel quel.

    Lève FileExistsError si le dossier existe déjà et n'est pas vide, sauf si
    `ecraser` : perdre silencieusement une simulation précédente serait
    difficile à rattraper.
    """
    def etape(texte):
        if progression is not None:
            progression(texte)

    nom = "".join(c for c in nom.strip() if c not in '\\/:*?"<>|').strip()
    if not nom:
        raise ValueError("Nom de projet vide ou composé uniquement de caractères interdits.")
    dossier = os.path.join(racine_results, nom)
    if os.path.isdir(dossier) and os.listdir(dossier) and not ecraser:
        raise FileExistsError(dossier)
    os.makedirs(dossier, exist_ok=True)
    if r_min is None:
        r_min = float(geom.rayon()[0])

    fichiers = {"dossier": dossier}

    etape("Fichiers de gravure (DXF)...")
    fichiers["dxf"] = dxf_io.exporter_dxf_paire(
        geom, os.path.join(dossier, f"{nom}.dxf"), n_points=n_points_dxf)

    etape("Animation du déploiement...")
    if frames_animation is not None:
        angles, frames = frames_animation
    else:
        angles, frames = calculer_frames_deploiement(geom, r_min, spec.amplitude_max,
                                                      n_frames=n_frames_gif)
    chemin_gif = os.path.join(dossier, f"{nom}.gif")
    sauvegarder_gif_deploiement(angles, frames, chemin_gif, fps=15, aller_retour=True)
    fichiers["gif"] = chemin_gif

    if bornes is not None and hyperparametres is not None:
        etape("Configuration...")
        chemin_json = os.path.join(dossier, f"{nom}.json")
        sauvegarder_configuration(chemin_json, spec, parametres, bornes, hyperparametres,
                                   reglage_echappement)
        fichiers["configuration"] = chemin_json

    fichiers["rapport"] = construire_rapport(
        os.path.join(dossier, f"{nom}.docx"), geom, spec, parametres,
        reglage_echappement=reglage_echappement, reglage_thermique=reglage_thermique,
        amplitudes_deg=amplitudes_deg, pas_balayage_deg=pas_balayage_deg,
        jeu_goupilles_um=jeu_goupilles_um, r_min=r_min, chemin_gif=chemin_gif,
        origine=origine, progression=progression)
    etape("Projet enregistré.")
    return fichiers
